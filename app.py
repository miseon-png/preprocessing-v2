import base64
import copy
import io
import re
import sqlite3
import traceback
from datetime import datetime

import gspread
from google.oauth2.service_account import Credentials
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
import pandas as pd
import streamlit as st
import streamlit.components.v1 as components

st.set_page_config(page_title="샐러드 생산일보 및 원료코드 System", layout="wide")

DB_FILE = "stock_master.db"

# ---------------------------------------------------------
# Google Spreadsheet API 연동
# ---------------------------------------------------------
@st.cache_resource
def get_gspread_client():
    scopes = [
        "https://www.googleapis.com/auth/spreadsheets",
        "https://www.googleapis.com/auth/drive"
    ]
    gcp_secrets = st.secrets["gcp_service_account"]
    private_key = gcp_secrets.get("private_key", "").replace("\\n", "\n")
    
    creds_dict = {
        "type": "service_account",
        "project_id": gcp_secrets.get("project_id", ""),
        "private_key_id": gcp_secrets.get("private_key_id", ""),
        "private_key": private_key,
        "client_email": gcp_secrets.get("client_email", ""),
        "client_id": gcp_secrets.get("client_id", ""),
        "auth_uri": "https://accounts.google.com/o/oauth2/auth",
        "token_uri": "https://oauth2.googleapis.com/token",
        "auth_provider_x509_cert_url": "https://www.googleapis.com/oauth2/v1/certs",
        "client_x509_cert_url": gcp_secrets.get("client_x509_cert_url", "")
    }
    creds = Credentials.from_service_account_info(creds_dict, scopes=scopes)
    client = gspread.authorize(creds)
    sheet_url = "https://docs.google.com/spreadsheets/d/1WM472gdj_7HBUvJdLRgwIU0t7GAMnZ47GpsCJfB5pBo/edit"
    return client.open_by_url(sheet_url)

# ---------------------------------------------------------
# SQLite DB 마스터 관리 (이월 롯트 - Lock 방지 최적화 적용)
# ---------------------------------------------------------
def get_db_conn():
    # 다중 사용자 접속 시 database is locked 방지용 안전장치 추가
    return sqlite3.connect(DB_FILE, check_same_thread=False, timeout=10.0)

def setup_database():
    conn = get_db_conn()
    cur = conn.cursor()
    cur.execute("""
        CREATE TABLE IF NOT EXISTS closed_lots (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            item_name TEXT,
            yymmdd TEXT,
            avail_g REAL,
            vendor TEXT,
            closed_ym TEXT
        )
    """)
    conn.commit()
    conn.close()

def reset_database():
    conn = get_db_conn()
    cur = conn.cursor()
    cur.execute("DELETE FROM closed_lots")
    conn.commit()
    conn.close()

def fetch_closed_lots():
    conn = get_db_conn()
    cur = conn.cursor()
    cur.execute("SELECT MAX(closed_ym) FROM closed_lots")
    latest_ym = cur.fetchone()[0]
    
    lots = {"양상추": [], "양배추": [], "적채": []}
    if latest_ym:
        cur.execute(
            "SELECT item_name, yymmdd, avail_g, vendor FROM closed_lots WHERE closed_ym = ?",
            (latest_ym,)
        )
        for item_name, yymmdd, avail_g, vendor in cur.fetchall():
            if item_name in lots:
                lots[item_name].append({
                    "yymmdd": yymmdd,
                    "avail_g": avail_g,
                    "avail": avail_g,
                    "vendor": vendor,
                    "label": f"{yymmdd[2:4]}/{yymmdd[4:]} 이월"
                })
    conn.close()
    return latest_ym, lots

def store_closed_lots(target_ym, lot_data):
    conn = get_db_conn()
    cur = conn.cursor()
    cur.execute("DELETE FROM closed_lots WHERE closed_ym = ?", (target_ym,))
    for item_name, lots in lot_data.items():
        for lot in lots:
            rem_g = lot.get("avail", lot.get("avail_g", 0.0))
            if rem_g > 0:
                cur.execute(
                    "INSERT INTO closed_lots (item_name, yymmdd, avail_g, vendor, closed_ym) VALUES (?, ?, ?, ?, ?)",
                    (item_name, lot["yymmdd"], rem_g, lot.get("vendor", "K01"), target_ym)
                )
    conn.commit()
    conn.close()

setup_database()

# ---------------------------------------------------------
# 출력 렌더링 헬퍼
# ---------------------------------------------------------
def create_print_button(html_data, label="🖨️ 바로 출력 / PDF"):
    encoded = base64.b64encode(html_data.encode("utf-8")).decode("utf-8")
    script_html = f"""
    <!DOCTYPE html>
    <html>
    <head>
    <style>
        body {{ margin: 0; padding: 0; font-family: sans-serif; }}
        .btn {{
            display: inline-flex; align-items: center; justify-content: center;
            padding: 0.25rem 0.75rem; border-radius: 0.5rem; min-height: 38px;
            width: 100%; cursor: pointer; font-size: 15px; box-sizing: border-box;
            background-color: #ffffff; color: #31333F; border: 1px solid #dcdcdc;
            transition: 0.2s;
        }}
        .btn:hover {{ border-color: #FF4B4B; color: #FF4B4B; }}
    </style>
    <script>
        function runPrint() {{
            try {{
                var decoded = decodeURIComponent(escape(window.atob("{encoded}")));
                var frame = document.createElement('iframe');
                frame.style.position = 'absolute';
                frame.style.width = '0px';
                frame.style.height = '0px';
                frame.style.border = 'none';
                document.body.appendChild(frame);
                var doc = frame.contentWindow.document;
                doc.open();
                doc.write(decoded);
                doc.close();
                frame.contentWindow.focus();
                setTimeout(function() {{ frame.contentWindow.print(); }}, 400);
            }} catch(e) {{
                alert("인쇄 초기화 중 오류가 발생했습니다.");
            }}
        }}
    </script>
    </head>
    <body>
        <button class="btn" onclick="runPrint()">{label}</button>
    </body>
    </html>
    """
    components.html(script_html, height=45)

def build_general_table_html(title, df):
    table_markup = df.to_html(index=False, classes="report-grid").replace('\\n', '<br>')
    return f"""
    <!DOCTYPE html>
    <html>
    <head>
    <meta charset="utf-8">
    <title>{title}</title>
    <style>
        body {{ font-family: '맑은 고딕', sans-serif; padding: 20px; }}
        .header {{ display: flex; justify-content: space-between; align-items: center; margin-bottom: 15px; }}
        .title {{ font-size: 18pt; font-weight: bold; color: #1F4E79; margin: 0; }}
        .approval {{ border-collapse: collapse; text-align: center; font-size: 9pt; }}
        .approval th, .approval td {{ border: 1px solid #444; padding: 4px 8px; }}
        .approval th {{ background: #f0f0f0; width: 45px; }}
        .approval td {{ height: 35px; width: 55px; }}
        .report-grid {{ width: 100%; border-collapse: collapse; font-size: 9pt; margin-top: 10px; }}
        .report-grid th, .report-grid td {{ border: 1px solid #444; padding: 5px; text-align: center; }}
        .report-grid th {{ background: #D9E1F2; color: #1F4E79; font-weight: bold; }}
        @media print {{
            body {{ padding: 0; }}
            @page {{ size: A4 portrait; margin: 12mm; }}
        }}
    </style>
    </head>
    <body>
        <div class="header">
            <h1 class="title">{title}</h1>
            <table class="approval">
                <tr><th rowspan="2">결<br>재</th><th>담당</th><th>검토</th><th>승인</th></tr>
                <tr><td></td><td></td><td></td></tr>
            </table>
        </div>
        {table_markup}
    </body>
    </html>
    """

def build_process_log_html(target_date, df_rows, is_filled=False):
    try:
        dt = datetime.strptime(target_date, "%Y-%m-%d")
        yyyy, mm, dd = dt.strftime("%Y"), dt.strftime("%m"), dt.strftime("%d")
    except:
        yyyy, mm, dd = "202 ", "  ", "  "

    rows_code = ""
    if df_rows.empty:
        rows_code = "<tr><td colspan='9' style='padding:20px; text-align:center;'>작업 내역이 없습니다.</td></tr>"
    else:
        current_product = None
        records = df_rows.to_dict("records")
        for rec in records:
            p_name = rec.get("제품명", "")
            item_name = rec.get("원재료", "")
            code_val = rec.get("통합_생성코드", "-")
            
            # 절단 규격 분류
            if "컷레터스" in p_name:
                cut_spec = "절단 6X6cm"
            elif any(k in item_name for k in ["양상추", "프릴", "로메인", "카이"]):
                cut_spec = "절단 4X4cm"
            elif any(k in item_name for k in ["양배추", "적채", "당근"]):
                cut_spec = "슬라이스 4mm"
            else:
                cut_spec = "절단 4X4cm"

            # 제품명 단위 rowspan 그룹화
            if p_name != current_product:
                current_product = p_name
                span_count = len([x for x in records if x.get("제품명") == p_name])
                formatted_p_name = str(p_name).replace("\n", "<br>")
                p_cell = f'<td rowspan="{span_count}" style="line-height:1.4;"><b>{formatted_p_name}</b></td>'
                check_cell = f'<td rowspan="{span_count}" style="font-size:8pt; text-align:left; line-height:1.4;">부패,변질,잔잎 등의 비가식부위 제거<br>원재료 상태 및 이물제거 확인</td>'
            else:
                p_cell = ""
                check_cell = ""

            post_val_td = f'<td style="text-align:right;">{rec.get("준비량(kg)", 0)} kg</td>' if is_filled else '<td style="text-align:right;">kg</td>'
            
            rows_code += f"""
            <tr>
                {p_cell}
                <td>{item_name}</td>
                <td>{code_val}</td>
                <td>적 / 부</td>
                {check_cell}
                <td>{cut_spec}</td>
                <td style="text-align:right;">{rec.get("투입량(kg)", 0)} kg</td>
                {post_val_td}
                <td></td>
            </tr>
            """

    for _ in range(3):
        rows_code += "<tr><td></td><td></td><td></td><td></td><td></td><td></td><td></td><td></td><td></td></tr>"

    return f"""
    <!DOCTYPE html>
    <html>
    <head>
    <meta charset="utf-8">
    <title>전처리 공정일지</title>
    <style>
        body {{ font-family: '맑은 고딕', sans-serif; padding: 15px; color: #000; background: #fff; margin: auto; max-width: 1200px; }}
        @media print {{
            body {{ padding: 0; max-width: none; }}
            @page {{ size: A4 landscape; margin: 12mm; }}
        }}
        .top-table {{ width: 100%; margin-bottom: 8px; }}
        .sheet-title {{ font-size: 24pt; font-weight: bold; letter-spacing: 2px; }}
        .approval {{ border-collapse: collapse; text-align: center; font-size: 9pt; float: right; }}
        .approval th, .approval td {{ border: 1px solid #000; padding: 4px 6px; }}
        .approval th {{ width: 45px; }}
        .approval td {{ height: 35px; width: 55px; }}
        .main-grid {{ width: 100%; border-collapse: collapse; font-size: 8.5pt; text-align: center; border: 1px solid #000; }}
        .main-grid th, .main-grid td {{ border: 1px solid #000; padding: 5px; }}
        .bg-light {{ background-color: #f2f2f2; font-weight: bold; }}
    </style>
    </head>
    <body>
        <table class="top-table">
            <tr>
                <td class="sheet-title">전처리 공정일지</td>
                <td>
                    <table class="approval">
                        <tr><th rowspan="2">결<br>재</th><th>작 성</th><th>검 토</th><th>승 인</th></tr>
                        <tr><td></td><td></td><td></td></tr>
                    </table>
                </td>
            </tr>
        </table>
        <table class="main-grid">
            <tr class="bg-light">
                <td colspan="2">작성일자</td>
                <td colspan="2" style="background: #fff;">{yyyy} 년 &nbsp; {mm} 월 &nbsp; {dd} 일</td>
                <td colspan="2" style="background: #fff;">&nbsp; : &nbsp; ~ &nbsp; : &nbsp;</td>
                <td>작성자</td>
                <td colspan="2" style="background: #fff;"></td>
            </tr>
            <tr>
                <td colspan="2" rowspan="2" class="bg-light">원료코드<br>부여방법</td>
                <td colspan="2" rowspan="2">원료 구분+입고일자 - 공급업체 구분</td>
                <td colspan="5" style="text-align:left; font-size:8pt; padding:3px 6px;">
                    카이피라(C), 프릴아이스(F), 버터헤드(B), 레드오크(O), 로메인(R), 케일(K), 적근대(S)<br>
                    치커리(E), 적채(A), 양배추(D), 양상추(X), 당근(G)
                </td>
            </tr>
            <tr>
                <td colspan="5" style="text-align:left; font-size:8pt; padding:3px 6px; border-top:1px dashed #aaa;">
                    공급업체 구분: 구름(G01), 에상스팜(H02), 승승장구(S01), 한스(K01)
                </td>
            </tr>
            <tr class="bg-light">
                <th rowspan="2">제품명</th>
                <th rowspan="2">원료명</th>
                <th colspan="7">작업원물 정보</th>
            </tr>
            <tr class="bg-light">
                <th>원료코드</th>
                <th>상태</th>
                <th>점검 항목</th>
                <th>절단규격</th>
                <th>생산량(전)</th>
                <th>투입량(후)</th>
                <th>비고</th>
            </tr>
            {rows_code}
        </table>
    </body>
    </html>
    """

# ---------------------------------------------------------
# 날짜 및 텍스트 파싱
# ---------------------------------------------------------
def extract_dates(raw):
    if pd.isna(raw):
        now = datetime.now()
        return now.strftime("%y"), now.strftime("%y%m%d"), now.strftime("%m%d")
    if isinstance(raw, (pd.Timestamp, datetime)):
        return raw.strftime("%y"), raw.strftime("%y%m%d"), raw.strftime("%m%d")
    
    txt = str(raw).strip()
    m_dot = re.search(r"(\d{2})\.(\d{2})(\d{2})", txt)
    if m_dot:
        yy, mm, dd = m_dot.group(1), m_dot.group(2), m_dot.group(3)
        return yy, f"{yy}{mm}{dd}", f"{mm}{dd}"
    
    m_hyphen = re.search(r"(\d{2,4})[-./](\d{1,2})[-./](\d{1,2})", txt)
    if m_hyphen:
        y_val = m_hyphen.group(1)
        yy = y_val[2:] if len(y_val) == 4 else y_val.zfill(2)
        mm = m_hyphen.group(2).zfill(2)
        dd = m_hyphen.group(3).zfill(2)
        return yy, f"{yy}{mm}{dd}", f"{mm}{dd}"
    
    m_six = re.search(r"\b(\d{6})\b", txt)
    if m_six:
        d = m_six.group(1)
        return d[:2], d, d[2:]
        
    now = datetime.now()
    return now.strftime("%y"), now.strftime("%y%m%d"), now.strftime("%m%d")

def identify_stock_sheet(xls):
    for s in xls.sheet_names:
        if s.strip().lower() == "sheet1":
            return s
    for s in xls.sheet_names:
        try:
            head = pd.read_excel(xls, sheet_name=s, header=None, nrows=10)
            merged_txt = " ".join(head.dropna().astype(str).values.flatten())
            if "양상추 매입" in merged_txt and "양배추 매입" in merged_txt:
                return s
        except:
            continue
    return xls.sheet_names[0]

def load_purchase_dataset(file_obj):
    if file_obj is None:
        return pd.DataFrame()
    fname = file_obj.name.lower()
    if fname.endswith((".xlsx", ".xls")):
        try:
            sheets = pd.read_excel(file_obj, sheet_name=None)
            return pd.concat(sheets.values(), ignore_index=True)
        except:
            return pd.DataFrame()
    for enc in ["utf-8-sig", "cp949", "euc-kr", "utf-8"]:
        try:
            file_obj.seek(0)
            return pd.read_csv(file_obj, encoding=enc)
        except:
            continue
    return pd.DataFrame()

# ---------------------------------------------------------
# [핵심] 롯트 컬럼 무관 블록 파싱 (순수 투입량 추출)
# ---------------------------------------------------------
KNOWN_INGREDIENTS = ["양상추", "양배추", "적채", "프릴", "로메인", "케일", "적근대", "치커리", "당근", "카이", "버터헤드", "레드오크"]

def parse_production_sheets(xls, excluded):
    sheets = [s for s in xls.sheet_names if s not in excluded and not s.startswith("재고") and not re.match(r"^\d{2}\.\d{2}$", s.strip())]
    results = []

    for s_name in sheets:
        df_sheet = pd.read_excel(xls, sheet_name=s_name, header=None)
        row_limit = len(df_sheet)

        for r_idx, row in df_sheet.iterrows():
            row_vals = [str(x).strip() for x in row.tolist()]
            if "발주" in row_vals and "수량" in row_vals:
                # 1. 일자 및 거래처(B열)
                _, yymmdd, _ = extract_dates(row_vals[0])
                record_date = f"20{yymmdd[:2]}-{yymmdd[2:4]}-{yymmdd[4:]}"
                vendor_title = row_vals[1] if len(row_vals) > 1 and row_vals[1] != "nan" else ""

                # 2. 원료명 컬럼 위치 매핑 (롯트 칸 존재 여부 무관)
                column_mapping = []
                for c_idx in range(4, len(row_vals)):
                    val = row_vals[c_idx]
                    if val != "nan" and val:
                        matched = next((ing for ing in KNOWN_INGREDIENTS if ing in val), None)
                        if matched:
                            column_mapping.append({"ingredient": matched, "col_idx": c_idx})

                # 3. 제품명 (기준행 바로 다음 행)
                if r_idx + 1 >= row_limit:
                    continue
                product_val = str(df_sheet.iloc[r_idx + 1, 0]).strip()
                if product_val in ["nan", ""]:
                    continue

                # 4. 차단벽(Break)을 둔 준비량 및 실투입량 행 탐색
                prep_idx, act_idx = None, None
                for scan in range(r_idx + 2, min(r_idx + 22, row_limit)):
                    scan_vals = [str(x).replace(" ", "") for x in df_sheet.iloc[scan].tolist()]
                    
                    # 다른 제품 블록이 시작되면 즉시 종료 (차단벽 로직)
                    if "발주" in scan_vals and "수량" in scan_vals:
                        break
                    
                    if prep_idx is None and any("전처리준비" in x for x in scan_vals):
                        prep_idx = scan
                    if act_idx is None and any("실투입" in x for x in scan_vals):
                        act_idx = scan

                # 5. 수량 추출
                for item in column_mapping:
                    c_pos = item["col_idx"]
                    ing_key = item["ingredient"]

                    prep_g, act_g = 0.0, 0.0
                    if prep_idx is not None:
                        try:
                            prep_g = float(df_sheet.iloc[prep_idx, c_pos])
                        except:
                            pass
                    if act_idx is not None:
                        try:
                            act_g = float(df_sheet.iloc[act_idx, c_pos])
                        except:
                            pass

                    if prep_g > 0 or act_g > 0:
                        results.append({
                            "std_date": record_date,
                            "vendor": vendor_title,
                            "product": product_val,
                            "ingredient": ing_key,
                            "prep_g": prep_g,
                            "act_g": act_g
                        })

    return sorted(results, key=lambda x: x["std_date"])

# ---------------------------------------------------------
# 수불부 계산 및 선입선출(FIFO)
# ---------------------------------------------------------
def calculate_subul_fifo(xls, sheet_name, prod_records):
    df_raw = pd.read_excel(xls, sheet_name=sheet_name, header=None)

    def to_float(val):
        if pd.isna(val): return 0.0
        try:
            num = re.sub(r"[^\d.]", "", str(val))
            return float(num) if num else 0.0
        except: return 0.0

    grouped_days = {}
    ordered_days = []

    # Sheet1에서 입고량 및 기말재고 가져오기
    for _, row in df_raw.iterrows():
        if len(row) == 0 or pd.isna(row.iloc[0]):
            continue
        first_val = str(row.iloc[0]).strip()
        if not first_val or any(k in first_val for k in ["날짜", "입고", "매입"]):
            continue

        _, yymmdd, _ = extract_dates(first_val)
        dt_str = f"20{yymmdd[:2]}-{yymmdd[2:4]}-{yymmdd[4:]}"

        y_in = to_float(row.iloc[1]) if len(row) > 1 else 0.0
        y_stk = to_float(row.iloc[2]) if len(row) > 2 else 0.0
        b_in = to_float(row.iloc[4]) if len(row) > 4 else 0.0
        b_stk = to_float(row.iloc[5]) if len(row) > 5 else 0.0
        j_in = to_float(row.iloc[7]) if len(row) > 7 else 0.0
        j_stk = to_float(row.iloc[8]) if len(row) > 8 else 0.0

        if dt_str not in grouped_days:
            grouped_days[dt_str] = {
                "date": dt_str, "yymmdd": yymmdd,
                "y_in": 0.0, "y_use": 0.0, "y_stk": 0.0,
                "b_in": 0.0, "b_use": 0.0, "b_stk": 0.0,
                "j_in": 0.0, "j_use": 0.0, "j_stk": 0.0
            }
            ordered_days.append(dt_str)

        grouped_days[dt_str]["y_in"] += y_in
        grouped_days[dt_str]["b_in"] += b_in
        grouped_days[dt_str]["j_in"] += j_in
        if y_stk > 0: grouped_days[dt_str]["y_stk"] = y_stk
        if b_stk > 0: grouped_days[dt_str]["b_stk"] = b_stk
        if j_stk > 0: grouped_days[dt_str]["j_stk"] = j_stk

    # 생산 실투입량 합산
    for r in prod_records:
        d = r["std_date"]
        ing = r["ingredient"]
        act = r["act_g"]
        if d not in grouped_days:
            grouped_days[d] = {
                "date": d, "yymmdd": d.replace("-", "")[2:],
                "y_in": 0.0, "y_use": 0.0, "y_stk": 0.0,
                "b_in": 0.0, "b_use": 0.0, "b_stk": 0.0,
                "j_in": 0.0, "j_use": 0.0, "j_stk": 0.0
            }
            ordered_days.append(d)

        if ing == "양상추": grouped_days[d]["y_use"] += act
        elif ing == "양배추": grouped_days[d]["b_use"] += act
        elif ing == "적채": grouped_days[d]["j_use"] += act

    ordered_days.sort()

    _, prior_lots = fetch_closed_lots()
    queues = {
        "양상추": copy.deepcopy(prior_lots["양상추"]),
        "양배추": copy.deepcopy(prior_lots["양배추"]),
        "적채": copy.deepcopy(prior_lots["적채"])
    }

    subul_output = []
    subul_code_map = {}

    for d in ordered_days:
        info = grouped_days[d]
        ymd = info["yymmdd"]
        mmdd_tag = f"{ymd[2:4]}/{ymd[4:]}"

        if info["y_in"] > 0: queues["양상추"].append({"yymmdd": ymd, "label": f"{mmdd_tag} 매입", "avail": info["y_in"], "vendor": "K01"})
        if info["b_in"] > 0: queues["양배추"].append({"yymmdd": ymd, "label": f"{mmdd_tag} 매입", "avail": info["b_in"], "vendor": "K01"})
        if info["j_in"] > 0: queues["적채"].append({"yymmdd": ymd, "label": f"{mmdd_tag} 매입", "avail": info["j_in"], "vendor": "K01"})

        def apply_fifo(item_key, code_prefix, needed_g):
            if needed_g <= 0: return "-", "-"
            rem = needed_g
            src_list, code_list = [], []
            for lot in queues[item_key]:
                if rem <= 0: break
                avail = lot.get("avail", 0.0)
                if avail <= 0: continue
                take = min(avail, rem)
                lot["avail"] = avail - take
                rem -= take
                v_code = lot.get("vendor", "K01")
                src_list.append(f"{lot['label']}({round(take/1000.0, 1)}kg)")
                code_list.append(f"{code_prefix}{lot['yymmdd']}-{v_code}")
            if rem > 0:
                src_list.append(f"부족({round(rem/1000.0, 1)}kg)")
                code_list.append(f"{code_prefix}{ymd}-K01[부족]")
            return " / ".join(src_list), " / ".join(list(dict.fromkeys(code_list)))

        y_src, y_c = apply_fifo("양상추", "X", info["y_use"])
        b_src, b_c = apply_fifo("양배추", "D", info["b_use"])
        j_src, j_c = apply_fifo("적채", "A", info["j_use"])

        subul_code_map[(d, "양상추")] = {"code": y_c, "source": y_src}
        subul_code_map[(d, "양배추")] = {"code": b_c, "source": b_src}
        subul_code_map[(d, "적채")] = {"code": j_c, "source": j_src}

        is_off = (sum([info["y_in"], info["y_use"], info["b_in"], info["b_use"], info["j_in"], info["j_use"]]) == 0)
        subul_output.append({
            "date": d,
            "yangsangchu_in": round(info["y_in"]/1000.0, 1), "yangsangchu_use": round(info["y_use"]/1000.0, 1),
            "yangsangchu_source": y_src, "yangsangchu_code": y_c, "yangsangchu_stock": round(info["y_stk"]/1000.0, 1),
            "yangbaechu_in": round(info["b_in"]/1000.0, 1), "yangbaechu_use": round(info["b_use"]/1000.0, 1),
            "yangbaechu_source": b_src, "yangbaechu_code": b_c, "yangbaechu_stock": round(info["b_stk"]/1000.0, 1),
            "jeokchae_in": round(info["j_in"]/1000.0, 1), "jeokchae_use": round(info["j_use"]/1000.0, 1),
            "jeokchae_source": j_src, "jeokchae_code": j_c, "jeokchae_stock": round(info["j_stk"]/1000.0, 1),
            "note": "휴무" if is_off else ""
        })

    # Excel 워크북 생성
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "원재료 수불부"
    ws.views.sheetView[0].showGridLines = True
    ws.merge_cells("A1:Q1")
    ws["A1"] = "원재료 수불부 (메인 3종 - kg 단위)"
    ws["A1"].font = Font(name="맑은 고딕", size=15, bold=True, color="1F4E79")
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")

    fill_main = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    fill_sub = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
    font_main = Font(name="맑은 고딕", size=10, bold=True, color="FFFFFF")
    font_sub = Font(name="맑은 고딕", size=9, bold=True, color="1F4E79")
    border_thin = Border(left=Side(style="thin", color="D3D3D3"), right=Side(style="thin", color="D3D3D3"),
                         top=Side(style="thin", color="D3D3D3"), bottom=Side(style="thin", color="D3D3D3"))

    ws.merge_cells("A2:A3"); ws["A2"] = "일자"
    ws.merge_cells("B2:F2"); ws["B2"] = "양상추(kg)"
    ws.merge_cells("G2:K2"); ws["G2"] = "양배추(kg)"
    ws.merge_cells("L2:P2"); ws["L2"] = "적채(kg)"
    ws.merge_cells("Q2:Q3"); ws["Q2"] = "비고"

    for col in range(1, 18):
        c = ws.cell(row=2, column=col)
        c.fill = fill_main; c.font = font_main; c.alignment = Alignment(horizontal="center", vertical="center")

    sub_titles = ["기초재고", "입고량", "투입량", "소진 매입일자", "기말재고"]
    for offset, base_col in enumerate([2, 7, 12]):
        for idx, sub_t in enumerate(sub_titles):
            cell = ws.cell(row=3, column=base_col + idx, value=sub_t)
            cell.fill = fill_sub; cell.font = font_sub; cell.alignment = Alignment(horizontal="center", vertical="center")

    r_curr = 4
    for item in subul_output:
        ws.cell(row=r_curr, column=1, value=item["date"]).alignment = Alignment(horizontal="center", vertical="center")
        ws.cell(row=r_curr, column=17, value=item["note"]).alignment = Alignment(horizontal="left", vertical="center")

        col_groups = [
            (2, item["yangsangchu_in"], item["yangsangchu_use"], item["yangsangchu_source"], item["yangsangchu_stock"]),
            (7, item["yangbaechu_in"], item["yangbaechu_use"], item["yangbaechu_source"], item["yangbaechu_stock"]),
            (12, item["jeokchae_in"], item["jeokchae_use"], item["jeokchae_source"], item["jeokchae_stock"])
        ]
        for base, in_val, use_val, src_val, stk_val in col_groups:
            stk_col_letter = get_column_letter(base + 4)
            if r_curr == 4:
                ws.cell(row=r_curr, column=base, value=round(stk_val + use_val - in_val, 1))
            else:
                ws.cell(row=r_curr, column=base, value=f"={stk_col_letter}{r_curr-1}")

            ws.cell(row=r_curr, column=base + 1, value=in_val)
            ws.cell(row=r_curr, column=base + 2, value=use_val)
            ws.cell(row=r_curr, column=base + 3, value=src_val).alignment = Alignment(horizontal="center", vertical="center")
            ws.cell(row=r_curr, column=base + 4, value=stk_val)

        for col in range(1, 18):
            cl = ws.cell(row=r_curr, column=col)
            cl.font = Font(name="맑은 고딕", size=9); cl.border = border_thin
            if col in [2, 3, 4, 6, 7, 8, 9, 11, 12, 13, 14, 16]:
                cl.number_format = "#,##0.0"
                cl.alignment = Alignment(horizontal="right", vertical="center")
        r_curr += 1

    wb_buf = io.BytesIO()
    wb.save(wb_buf)

    # 마감 잔여 롯트 이월 계산
    last_day_record = subul_output[-1] if subul_output else {}
    target_stocks_g = {
        "양상추": last_day_record.get("yangsangchu_stock", 0.0) * 1000.0,
        "양배추": last_day_record.get("yangbaechu_stock", 0.0) * 1000.0,
        "적채": last_day_record.get("jeokchae_stock", 0.0) * 1000.0
    }
    rem_lots = {"양상추": [], "양배추": [], "적채": []}
    for k in ["양상추", "양배추", "적채"]:
        t_g = target_stocks_g[k]
        alive = [q for q in queues[k] if q.get("avail", 0.0) > 0]
        for q in reversed(alive):
            if t_g <= 0: break
            fit_g = min(q["avail"], t_g)
            rem_lots[k].insert(0, {
                "yymmdd": q["yymmdd"], "avail_g": fit_g, "avail": fit_g, "vendor": q.get("vendor", "K01")
            })
            t_g -= fit_g

    return wb_buf.getvalue(), subul_output, subul_code_map, rem_lots

# ---------------------------------------------------------
# 전체 원료 코드 생성기 (출고량/공정일지용)
# ---------------------------------------------------------
PREFIX_MAP = {
    "양상추": "X", "양배추": "D", "적채": "A", "프릴": "F", "로메인": "R",
    "케일": "K", "적근대": "S", "치커리": "E", "당근": "G", "카이": "C",
    "버터헤드": "B", "레드오크": "O"
}
VENDOR_MAP = {"한스": "K01", "승승장구": "S01", "에상스팜": "H02", "넥스토팜": "N01"}

def generate_unified_codes(records, subul_map, df_purchase):
    p_entries = []
    if not df_purchase.empty:
        col_dt = next((c for c in df_purchase.columns if any(k in str(c) for k in ["일자", "기간", "날짜"])), None)
        col_it = next((c for c in df_purchase.columns if any(k in str(c) for k in ["품명", "품목", "원재료"])), None)
        col_vn = next((c for c in df_purchase.columns if any(k in str(c) for k in ["거래처", "매입처", "공급"])), None)

        if col_dt and col_it:
            for _, row in df_purchase.iterrows():
                dt_raw, it_raw = row[col_dt], str(row[col_it]).strip()
                vn_raw = str(row[col_vn]).strip() if col_vn and pd.notna(row[col_vn]) else ""
                if pd.isna(dt_raw) or not it_raw or vn_raw == "기타":
                    continue
                _, p_yymmdd, _ = extract_dates(dt_raw)
                p_date = f"20{p_yymmdd[:2]}-{p_yymmdd[2:4]}-{p_yymmdd[4:]}"
                matched = next((k for k in PREFIX_MAP.keys() if k in it_raw), it_raw)
                pref = PREFIX_MAP.get(matched, "M")
                v_code = VENDOR_MAP.get(vn_raw, "K01")
                p_entries.append({
                    "date": p_date, "yymmdd": p_yymmdd, "item": matched,
                    "code": f"{pref}{p_yymmdd}-{v_code}"
                })

    process_rows = []
    summary_rows = []
    summary_checker = set()

    for r in records:
        d = r["std_date"]
        ing = r["ingredient"]
        v_title = r.get("vendor", "")
        p_name = f"{v_title}\n{r['product']}" if v_title else r["product"]

        pref = PREFIX_MAP.get(ing, "M")
        final_code = "-"
        src_note = "-"

        # 메인 3종: 수불부 FIFO 적용
        if ing in ["양상추", "양배추", "적채"]:
            data = subul_map.get((d, ing), {})
            final_code = data.get("code", "-")
            src_note = data.get("source", "-")
        # 서브 원료: 최근 매입일 역추적
        else:
            candidates = [p for p in p_entries if p["item"] == ing and p["date"] <= d]
            if candidates:
                candidates.sort(key=lambda x: x["date"], reverse=True)
                final_code = candidates[0]["code"]
                src_note = f"{candidates[0]['yymmdd'][2:4]}/{candidates[0]['yymmdd'][4:]} 매입 참조"

        if final_code == "-":
            final_code = f"{pref}{d.replace('-', '')[2:]}-K01(임시)"

        act_kg = round(r["act_g"] / 1000.0, 1)
        prep_kg = round(r["prep_g"] / 1000.0, 1)

        process_rows.append({
            "생산일자": d, "제품명": p_name, "원재료": ing,
            "통합_생성코드": final_code, "투입량(kg)": act_kg, "준비량(kg)": prep_kg
        })

        summary_key = (d, ing, final_code, src_note)
        if summary_key not in summary_checker and act_kg > 0:
            summary_checker.add(summary_key)
            summary_rows.append({
                "날짜": d, "원재료": ing, "생성코드": final_code,
                "투입량(kg)": act_kg, "준비량(kg)": prep_kg, "소진출처(상세)": src_note
            })

    return pd.DataFrame(process_rows), pd.DataFrame(summary_rows)

# ---------------------------------------------------------
# Streamlit UI
# ---------------------------------------------------------
st.title("🥗 샐러드 생산일보 원료코드 및 수불 관리 시스템")

with st.sidebar:
    st.header("🗄️ 수불 이월 DB 관리")
    last_ym, saved_lots = fetch_closed_lots()
    if last_ym:
        st.success(f"📌 저장된 마감월: **{last_ym}**")
        with st.expander("이월 잔여 롯트 확인"):
            for k, lots in saved_lots.items():
                if lots:
                    st.caption(f"**{k}**")
                    for l in lots:
                        st.write(f"- 입고일: {l['yymmdd']} | 잔여: {round(l['avail_g']/1000.0, 1)}kg")
    else:
        st.info("💡 등록된 이월 DB가 없습니다.\n당월 엑셀의 입고분부터 순서대로 소진됩니다.")

    st.markdown("---")
    if st.button("🗑️ 과거 이월 DB 초기화", type="secondary"):
        reset_database()
        st.success("✅ DB가 초기화되었습니다.")
        st.rerun()

col_up1, col_up2 = st.columns(2)
with col_up1:
    f_excel = st.file_uploader("1. 샐러드 생산일지 (Excel) 업로드", type=["xlsx", "xls"])
with col_up2:
    f_csv = st.file_uploader("2. 원료 매입자료 (CSV/Excel) 업로드", type=["csv", "xlsx", "xls"])

# 5개 탭 구성 (원료코드 대조표 탭 제외)
tab_subul, tab_codes, tab_close, tab_blank, tab_filled = st.tabs([
    "📊 메인 3종 원재료 수불부",
    "☁️ 전체 원료코드 및 구글시트 연동",
    "🔒 당월 수불 마감 확정",
    "📋 전처리 공정일지 (빈칸 출력)",
    "📋 전처리 공정일지 (완성본)"
])

if f_excel is not None:
    try:
        df_purchase = load_purchase_dataset(f_csv)
        xls_file = pd.ExcelFile(f_excel)
        stock_sheet = identify_stock_sheet(xls_file)
        ignore_sheets = ["테스트", "원가", stock_sheet, "Sheet1", "sheet1"]

        # 블록 파싱 (롯트 칸 유무 무관)
        prod_records = parse_production_sheets(xls_file, ignore_sheets)
        if not prod_records:
            st.error("생산일지 시트에서 투입량이 존재하는 제품 블록을 찾지 못했습니다.")
            st.stop()

        # 수불 계산 & 코드 통합
        excel_bytes, subul_rows, fifo_map, end_lots = calculate_subul_fifo(xls_file, stock_sheet, prod_records)
        df_process, df_summary = generate_unified_codes(prod_records, fifo_map, df_purchase)

        # -----------------------------------------------------
        # Tab 1: 메인 3종 원재료 수불부
        # -----------------------------------------------------
        with tab_subul:
            st.subheader("📋 메인 3종 원재료 수불부 (선입선출 적용)")
            df_subul = pd.DataFrame(subul_rows)
            cols_order = [
                "date", "yangsangchu_in", "yangsangchu_use", "yangsangchu_source", "yangsangchu_stock",
                "yangbaechu_in", "yangbaechu_use", "yangbaechu_source", "yangbaechu_stock",
                "jeokchae_in", "jeokchae_use", "jeokchae_source", "jeokchae_stock", "note"
            ]
            df_subul = df_subul[[c for c in cols_order if c in df_subul.columns]].rename(columns={
                "date": "일자",
                "yangsangchu_in": "양상추_입고", "yangsangchu_use": "양상추_투입", "yangsangchu_source": "양상추_소진출처", "yangsangchu_stock": "양상추_재고",
                "yangbaechu_in": "양배추_입고", "yangbaechu_use": "양배추_투입", "yangbaechu_source": "양배추_소진출처", "yangbaechu_stock": "양배추_재고",
                "jeokchae_in": "적채_입고", "jeokchae_use": "적채_투입", "jeokchae_source": "적채_소진출처", "jeokchae_stock": "적채_재고",
                "note": "비고"
            })

            col_btn1, col_btn2 = st.columns([3, 1])
            with col_btn1:
                st.download_button("📥 수불부 엑셀 다운로드 (.xlsx)", data=excel_bytes, file_name="메인3종_원재료_수불부.xlsx", use_container_width=True)
            with col_btn2:
                html_subul = build_general_table_html("메인 3종 원재료 수불부", df_subul)
                create_print_button(html_subul, "🖨️ 수불부 인쇄 / PDF")
            st.dataframe(df_subul, use_container_width=True)

        # -----------------------------------------------------
        # Tab 2: 전체 원료코드 및 구글시트 연동
        # -----------------------------------------------------
        with tab_codes:
            st.subheader("🏷️ 일자별 전체 원료코드 요약 및 구글 시트 연동")
            st.caption("양상추/양배추/적채는 선입선출, 서브 원료는 최근 매입일 기준으로 생성된 최종 코드 목록입니다.")
            
            if not df_summary.empty:
                sort_order = {"양상추": 1, "양배추": 2, "적채": 3, "프릴": 4}
                df_summary_disp = df_summary.copy()  # 경고 방지용 깊은 복사
                df_summary_disp["_sort"] = df_summary_disp["원재료"].map(lambda x: sort_order.get(x, 99))
                df_summary_disp = df_summary_disp.sort_values(by=["날짜", "_sort"]).drop(columns=["_sort"])

                col_b1, col_b2 = st.columns([3, 1])
                with col_b1:
                    code_buf = io.BytesIO()
                    with pd.ExcelWriter(code_buf, engine="openpyxl") as wr:
                        df_summary_disp.to_excel(wr, index=False, sheet_name="일자별_원료코드")
                    st.download_button("📥 전체 원료코드 목록 다운로드 (.xlsx)", data=code_buf.getvalue(), file_name="전체_원료코드_목록.xlsx", use_container_width=True)
                with col_b2:
                    html_codes = build_general_table_html("전체 원료코드 목록", df_summary_disp)
                    create_print_button(html_codes, "🖨️ 인쇄 / PDF 저장")
                
                st.dataframe(df_summary_disp, use_container_width=True)

                st.markdown("---")
                st.subheader("☁️ 구글 스프레드시트 출고량 일괄 연동")
                col_d_start, col_d_end = st.columns(2)
                with col_d_start:
                    date_from = st.date_input("동기화 시작일", value=datetime.today().replace(day=1), key="sync_from")
                with col_d_end:
                    date_to = st.date_input("동기화 종료일", value=datetime.today(), key="sync_to")

                if st.button("🚀 선택 기간 전체 원료 출고량 구글 시트 전송", type="primary", use_container_width=True):
                    with st.spinner("구글 스프레드시트 반영 중... (API 최적화 방식)"):
                        try:
                            s_str = date_from.strftime("%Y-%m-%d")
                            e_str = date_to.strftime("%Y-%m-%d")
                            df_target = df_summary_disp[(df_summary_disp["날짜"] >= s_str) & (df_summary_disp["날짜"] <= e_str)]

                            doc = get_gspread_client()
                            wks = doc.worksheet("시트1")
                            
                            # 구글 시트 전체 데이터를 한 번에 가져옵니다.
                            sheet_rows = wks.get_all_values()
                            header = sheet_rows[0] if sheet_rows else []
                            
                            # 유지할 데이터만 걸러냅니다.
                            retained_data = [header] if header else []
                            deleted_count = 0
                            
                            for row_val in sheet_rows[1:]:
                                try:
                                    r_dt = row_val[0].strip()
                                    # 기간 내의 "시스템자동" 데이터라면 삭제 대상이므로 보존 리스트에서 제외
                                    if s_str <= r_dt <= e_str and len(row_val) > 2 and row_val[2].strip() == "시스템자동":
                                        deleted_count += 1
                                        continue 
                                    retained_data.append(row_val)
                                except:
                                    retained_data.append(row_val) # 예외 발생 시 안전하게 보존

                            # 새로 전송할 데이터를 생성합니다.
                            new_payload = []
                            for _, rec in df_target.iterrows():
                                if rec["투입량(kg)"] > 0:
                                    new_payload.append([
                                        rec["날짜"], "출고", "시스템자동", rec["원재료"],
                                        rec["투입량(kg)"], 0, 0, rec["생성코드"]
                                    ])

                            # 보존 데이터와 새 데이터를 병합합니다.
                            final_data = retained_data + new_payload

                            # 구글 시트를 비우고 통째로 덮어씌웁니다. (API 429 에러 완벽 차단)
                            wks.clear()
                            if final_data:
                                wks.update(values=final_data, range_name="A1")
                                st.success(f"✅ 기존 데이터 {deleted_count}건 삭제 및 신규 데이터 {len(new_payload)}건 등록(덮어쓰기)을 완료했습니다!")
                                st.balloons()
                            else:
                                st.success("✅ 기존 데이터를 정리했습니다. 전송할 신규 데이터가 없습니다.")
                        except Exception as ex:
                            st.error(f"구글 시트 연동 실패: {ex}")

        # -----------------------------------------------------
        # Tab 3: 당월 수불 마감 확정
        # -----------------------------------------------------
        with tab_close:
            st.subheader("🔒 당월 수불 마감 확정")
            default_close_ym = subul_rows[-1]["date"][:7] if subul_rows else datetime.now().strftime("%Y-%m")
            input_ym = st.text_input("마감 확정 연월 (YYYY-MM)", value=default_close_ym)

            st.write("▼ **익월로 이월될 메인 3종 잔여 롯트 현황**")
            for item_k, l_list in end_lots.items():
                total_kg = sum(x["avail_g"] for x in l_list) / 1000.0
                st.write(f"• **{item_k}**: 총 {round(total_kg, 1)}kg (잔여 롯트 {len(l_list)}건)")
                for l in l_list:
                    st.caption(f"   - 입고일: {l['yymmdd']} | 잔여량: {round(l['avail_g']/1000.0, 1)}kg")

            if st.button("🔒 당월 수불 마감 확정 및 DB 저장", type="primary"):
                store_closed_lots(input_ym, end_lots)
                st.balloons()
                st.success(f"✅ **{input_ym}** 마감이 성공적으로 저장되었습니다.")

        # -----------------------------------------------------
        # Tab 4 & 5: 전처리 공정일지
        # -----------------------------------------------------
        for tab_target, is_done in zip([tab_blank, tab_filled], [False, True]):
            with tab_target:
                type_name = "완성본" if is_done else "빈칸 출력용"
                st.subheader(f"📋 전처리 공정일지 ({type_name})")
                if not df_process.empty:
                    available_dates = sorted(df_process["생산일자"].unique())
                    col_sel, _ = st.columns([1, 2])
                    with col_sel:
                        selected_date = st.selectbox(f"📅 출력 일자 선택 ({type_name})", available_dates, key=f"sel_dt_{is_done}")

                    df_sheet_target = df_process[df_process["생산일자"] == selected_date].copy()
                    item_order = {"양상추": 1, "양배추": 2, "적채": 3, "프릴": 4}
                    df_sheet_target["_order"] = df_sheet_target["원재료"].map(lambda x: item_order.get(x, 99))
                    df_sheet_target = df_sheet_target.sort_values(by=["제품명", "_order"])

                    html_log = build_process_log_html(selected_date, df_sheet_target, is_filled=is_done)

                    col_p_btn, _ = st.columns([1, 3])
                    with col_p_btn:
                        create_print_button(html_log, f"🖨️ 공정일지 {type_name} 바로 출력")
                    components.html(html_log, height=750, scrolling=True)

    except Exception as err:
        st.error(f"실행 중 예외가 발생했습니다: {err}")
        st.code(traceback.format_exc())
else:
    for t in [tab_subul, tab_codes, tab_close, tab_blank, tab_filled]:
        with t:
            st.info("👆 상단에서 1. 샐러드 생산일지(Excel)와 2. 원료 매입자료(CSV/Excel)를 업로드해 주세요.")
